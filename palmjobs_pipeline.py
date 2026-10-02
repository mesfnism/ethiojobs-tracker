"""
Palm Jobs Pipeline — fourth job board (palmjobs.et), a Next.js app backed
by a public Supabase Postgres table. Unlike every other site in this
project, this one is NOT scraped by parsing rendered text: its `jobs`
table is readable anonymously (confirmed while signed out of any account —
browsing it never touches personal/candidate data), and openly serves
clean structured fields most other sites only state in free text:
min/max salary with currency (ETB or USD), education level, field of
study, required skills (as an actual list), experience level, and more.

How this gets the data, and why: `<project>.supabase.co/rest/v1/jobs` is
the exact same public endpoint the site's own front end uses to render
every job, with the public anon key every visitor's browser already
downloads as part of the page's JS bundle. This pipeline fetches the
/jobs page's HTML, finds that anon key by reading the same JS bundle
(see discover_api_credentials below), and reuses it to page through the
SAME table itself, in the SAME shape — not a different, non-public API,
not a different request shape, not a bypassed protection of any kind. It
never touches account-scoped tables (candidates, resumes, notifications,
user_*) — only the public `jobs` and `companies_public` tables the site
already shows every anonymous visitor. (This used to run through a real
Playwright browser instead of plain HTTP; see the 2026-10-02 updates
below for why that was dropped.)

Job IDs come directly from the API's own `id` field (a real UUID), so
dedup across runs is exact — no derived/hashed identifier needed, unlike
sites where only scraped text is available.

Update (2026-10-02, part 1): a real run on GitHub Actions reported "Could
not find API credentials by either method," while the exact same code,
run locally minutes later, found them immediately — confirming the same
automated-browser-detection pattern already found for reporterjobs_pipeline.py
(see that file's docstring). That same local run also surfaced two
further, unrelated, genuine bugs, both fixed in the same pass: (a) main()
captured `page.context.request` then called `page.close()` on a page
created via `browser.new_page()`, whose IMPLICIT context is tied 1:1 to
that page — closing the page tore down the context the APIRequestContext
depended on, breaking every subsequent call with "Target page, context or
browser has been closed"; and (b) a captured live request's FULL header
set was being reused verbatim for a different query than the one it was
captured from, which broke pagination with PostgREST's PGRST116 error
("Cannot coerce the result to a single JSON object") whenever the
captured request happened to be a different (singular-result) query than
our own list query.

Update (2026-10-02, part 2): a second GitHub Actions run, after fixing
both of those bugs, STILL failed with "Could not find API credentials,"
while a local run with the identical code succeeded and pulled 468
postings cleanly. That confirms this is the same automation-fingerprint
blocking as ReporterJobs, not a flaky one-off — so this file now drops
Playwright entirely, the same fix applied there: curl_cffi (TLS-fingerprint
-impersonating plain HTTP, no browser, no CDP) fetches the /jobs page's
HTML and its JS chunks to find the anon key, and also fetches the
paginated REST API data itself — which was always a plain authenticated
REST call, never something that actually needed a real browser. The
capture_api_headers() "intercept a live request" method is dropped along
with it: that only ever worked by watching a real client-side JS XHR
fire, which has no equivalent when nothing executes any JavaScript at
all; discover_api_credentials() (reading the public anon key straight out
of the JS bundle text) was always the dominant, more reliable path
anyway, so nothing of real value is lost.

Update (2026-10-02, part 3): the first GitHub Actions run of the
curl_cffi rewrite got back a 429 (rate limited) loading LIST_PAGE_URL,
not a 403. That's a meaningfully different signal from ReporterJobs'
dead-end 403s: a 429 means the site understood and accepted the request
as a legitimate-looking client, just one going too fast, not one it has
fingerprinted as automation. Added _get_with_retry() — backs off and
retries on 429 (honoring a Retry-After header if the site sends one),
used for every GET this file makes (the /jobs page, each JS chunk, and
the paginated REST calls). Left genuinely unresolved: whether that 429
was PalmJobs' own rate limiter or a side effect of GitHub Actions runners
sharing IP ranges with other traffic — the retry logic handles either
cause the same way, so it isn't necessary to tell them apart to move
forward.
"""

import os
import re
import sys
import time
from pathlib import Path
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from bs4 import BeautifulSoup

try:
    from curl_cffi import requests as curl_requests
except ImportError:
    curl_requests = None

# TLS-fingerprint to impersonate (see this file's 2026-10-02 part-2 docstring
# note) — "chrome124" matches reporterjobs_pipeline.py's own default, kept
# overridable via env var the same way, in case one site's WAF ever reacts
# differently to a particular Chrome version's fingerprint than the other's.
IMPERSONATE = os.environ.get("PALMJOBS_IMPERSONATE", "chrome124")

LIST_PAGE_URL = "https://palmjobs.et/jobs"
_CHUNK_SRC_RE = re.compile(r"/_next/static/chunks/")

PAGE_SIZE = int(os.environ.get("PALMJOBS_PAGE_SIZE", "50"))
MAX_PAGES = int(os.environ.get("PALMJOBS_MAX_PAGES", "40"))
STOP_AFTER_CONSECUTIVE_SEEN = int(os.environ.get("PALMJOBS_STOP_AFTER_SEEN", "25"))
REQUEST_DELAY_SECONDS = float(os.environ.get("PALMJOBS_REQUEST_DELAY", "1.0"))

OUTPUT_XLSX = os.environ.get("PALMJOBS_OUTPUT_XLSX", "palmjobs_tracker.xlsx")

SELECT_FIELDS = (
    "id,employer_id,company_id,company_name,job_title,job_location,"
    "open_positions,job_type,experience_level,min_salary,max_salary,"
    "currency,job_description,application_deadline,external_link,"
    "job_status,date_posted,created_at,job_industry,education_level,"
    "field_of_study,required_skills,is_aggregated,source_url,source_name,"
    "email_application"
)

COLUMNS = [
    "job_id", "job_title", "employer", "category", "location", "work_type",
    "career_level", "employment_type", "number_required",
    "education_required", "years_experience_required", "skills_required",
    "special_skill_training", "salary_hint", "application_deadline",
    "how_to_apply", "source_url", "source", "date_posted_relative",
    "date_scraped", "extraction_method", "description",
]

NAVY = "1F3B57"

_EDUCATION_NORMALIZE = [
    (re.compile(r"\bph\.?d\b|\bdoctorate\b|\bdoctoral\b", re.I), "PhD"),
    (re.compile(r"\bmaster'?s?\b", re.I), "Master's degree"),
    (re.compile(r"\bbachelor'?s?\b|\bundergraduate\b", re.I), "Bachelor's degree"),
    (re.compile(r"\bdiploma\b", re.I), "Diploma"),
    (re.compile(r"\bcertificate\b|\btvet\b", re.I), "Certificate"),
]


def normalize_education(raw):
    if not raw:
        return None
    for pattern, label in _EDUCATION_NORMALIZE:
        if pattern.search(raw):
            return label
    return None  # left unclassified rather than guessed — matches ISCED_LEVELS keys only


def build_salary_hint(row):
    """Builds a text salary_hint from the API's structured min/max/currency
    fields, in the same shape taxonomy.salary_bucket already parses
    ("<currency> <amount>" or "<currency> <lo> - <hi>"), so the existing
    salary chart picks these up with no special-casing — except that
    taxonomy.salary_bucket only buckets Birr/ETB figures; a USD salary
    here is preserved as text but won't be bucketed (honest gap, not a
    silent currency conversion)."""
    currency = (row.get("currency") or "").strip()
    lo, hi = row.get("min_salary"), row.get("max_salary")
    if lo is None and hi is None:
        return None
    if lo is not None and hi is not None and lo != hi:
        return f"{currency} {lo} - {hi}".strip()
    amount = lo if lo is not None else hi
    return f"{currency} {amount}".strip()


def structure_job(row):
    skills = row.get("required_skills")
    if isinstance(skills, list):
        skills_required = "; ".join(str(s) for s in skills if s)
    else:
        skills_required = skills

    return {
        "job_id": row.get("id"),
        "job_title": row.get("job_title"),
        "employer": row.get("company_name"),
        "category": row.get("job_industry"),
        "location": row.get("job_location"),
        "work_type": None,
        "career_level": None,
        "employment_type": row.get("job_type"),
        "number_required": row.get("open_positions"),
        "education_required": normalize_education(row.get("education_level")),
        "years_experience_required": row.get("experience_level"),
        "skills_required": skills_required or None,
        "special_skill_training": None,
        "salary_hint": build_salary_hint(row),
        "application_deadline": row.get("application_deadline"),
        "how_to_apply": row.get("external_link") or row.get("email_application"),
        "source_url": f"https://palmjobs.et/jobs?jobId={row.get('id')}",
        "source": "PalmJobs",
        "date_posted_relative": row.get("date_posted"),
        "date_scraped": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"),
        "extraction_method": "public_api",
        # job_description doubles as the free-text source for education
        # specialization extraction (field_of_study is also returned
        # directly by the API and could be used as-is in a future pass).
        "description": row.get("job_description") or row.get("field_of_study"),
    }


# --------------------------------------------------------------------------
# Plain-HTTP (curl_cffi) credential discovery and paginated data fetch
# --------------------------------------------------------------------------

_ANON_KEY_PATTERN = r"eyJ[a-zA-Z0-9_\-\.]{60,}"
_SUPABASE_URL_PATTERN = r"https://[a-z0-9]+\.supabase\.co"

RETRY_429_MAX_ATTEMPTS = int(os.environ.get("PALMJOBS_429_MAX_RETRIES", "4"))
RETRY_429_BACKOFF_SECONDS = float(os.environ.get("PALMJOBS_429_BACKOFF", "8"))


def _get_with_retry(session, url, headers=None, timeout=30):
    """GETs a URL, retrying with backoff on HTTP 429 (rate limited).

    Added 2026-10-02 (part 3): the very first GitHub Actions run of the
    curl_cffi rewrite got a 429 loading LIST_PAGE_URL — a genuinely
    different signal from the 403s documented elsewhere in this file and
    in reporterjobs_pipeline.py's docstring. A 403 from those sites means
    "you're detected as automation, full stop" — retrying changes
    nothing. A 429 means "you're a legitimate-looking client going too
    fast" — the site accepted and understood the request well enough to
    rate-limit it rather than block it outright, so backing off and
    retrying is the right response, not a reason to declare this blocked
    too. Honors a Retry-After header when the server sends one; otherwise
    waits an increasing backoff. Still returns the (non-retryable) response
    as-is on any other status, or after exhausting retries, so callers'
    existing status/exception handling is unchanged."""
    attempt = 0
    while True:
        resp = session.get(url, headers=headers, timeout=timeout)
        if resp.status_code != 429 or attempt >= RETRY_429_MAX_ATTEMPTS:
            return resp
        retry_after = None
        try:
            retry_after = resp.headers.get("Retry-After")
        except Exception:
            pass
        try:
            wait = float(retry_after) if retry_after else RETRY_429_BACKOFF_SECONDS * (attempt + 1)
        except (TypeError, ValueError):
            wait = RETRY_429_BACKOFF_SECONDS * (attempt + 1)
        print(f"  Got 429 (rate limited) fetching {url} — waiting {wait:.0f}s "
              f"before retry {attempt + 1}/{RETRY_429_MAX_ATTEMPTS}...")
        time.sleep(wait)
        attempt += 1


def _chunk_script_urls(html, base_url):
    """Finds every /_next/static/chunks/ script src on the page, resolved
    to an absolute URL. The chunk file that happens to hold the anon key
    is content-hashed and changes on every deployment, so this doesn't
    guess a filename — it returns every candidate chunk and
    discover_api_credentials below checks them all."""
    soup = BeautifulSoup(html, "html.parser")
    urls = []
    for tag in soup.find_all("script", src=True):
        src = tag["src"]
        if not _CHUNK_SRC_RE.search(src):
            continue
        if src.startswith("http"):
            urls.append(src)
        elif src.startswith("/"):
            urls.append("https://palmjobs.et" + src)
        else:
            urls.append(base_url.rstrip("/") + "/" + src)
    return urls


def discover_api_credentials(session):
    """Finds the site's own public Supabase anon key and project URL by
    reading its already-loaded JS bundle — the same public key and URL
    every visitor's browser already downloads just to render the page,
    not a different or bypassed credential.

    Update (2026-10-02, part 2): this used to run inside a real Playwright
    page via page.evaluate(), fetching each chunk with the browser's own
    fetch(). Rewritten here to do the exact same two steps as plain HTTP
    through curl_cffi instead: GET the /jobs page's HTML, find its chunk
    script tags with BeautifulSoup, then GET each chunk directly — no
    browser, no CDP, nothing for automation-fingerprint detection to catch.
    capture_api_headers() (watching for a live client-side XHR to
    /rest/v1/jobs) is dropped entirely along with Playwright: there is no
    JS execution here to watch, and that method was already the weaker,
    fallback path even when a real browser was available."""
    try:
        resp = _get_with_retry(session, LIST_PAGE_URL, timeout=30)
    except Exception as e:
        print(f"  Could not load {LIST_PAGE_URL}: {e}")
        return None, None
    if resp.status_code != 200:
        print(f"  Unexpected status {resp.status_code} loading {LIST_PAGE_URL}")
        return None, None

    chunk_urls = _chunk_script_urls(resp.text, LIST_PAGE_URL)
    anon_key, supa_url = None, None
    for chunk_url in chunk_urls:
        try:
            chunk_resp = _get_with_retry(session, chunk_url, timeout=30)
        except Exception:
            continue
        if chunk_resp.status_code != 200:
            continue
        text = chunk_resp.text
        if not anon_key:
            m = re.search(_ANON_KEY_PATTERN, text)
            if m:
                anon_key = m.group(0)
        if not supa_url:
            m2 = re.search(_SUPABASE_URL_PATTERN, text)
            if m2:
                supa_url = m2.group(0)
        if anon_key and supa_url:
            break

    if not anon_key or not supa_url:
        return None, None
    headers = {"apikey": anon_key, "Authorization": f"Bearer {anon_key}"}
    base_url = f"{supa_url}/rest/v1/jobs"
    return headers, base_url


def fetch_jobs_page(session, base_url, headers, offset):
    url = (
        f"{base_url}?select={SELECT_FIELDS}&job_status=eq.Active"
        f"&order=created_at.desc&limit={PAGE_SIZE}&offset={offset}"
    )
    # Explicit Accept: application/json for OUR OWN list query, regardless
    # of what's in `headers` — defense in depth against the 2026-10-02
    # PGRST116 bug (a reused Accept header from a different, singular-
    # result query once broke pagination with "Cannot coerce the result
    # to a single JSON object"): this request's own intent always wins.
    request_headers = {**headers, "Accept": "application/json"}
    resp = _get_with_retry(session, url, headers=request_headers, timeout=30)
    if resp.status_code != 200:
        raise RuntimeError(f"Unexpected status {resp.status_code} fetching jobs page: {resp.text[:300]}")
    return resp.json()


# --------------------------------------------------------------------------
# Excel read/write — identical pattern to the other pipelines
# --------------------------------------------------------------------------

def load_existing_ids(path: Path):
    if not path.exists():
        return set()
    wb = openpyxl.load_workbook(path, read_only=True)
    if "Jobs" not in wb.sheetnames:
        return set()
    ws = wb["Jobs"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    id_col = header.index("job_id")
    ids = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[id_col]:
            ids.add(str(row[id_col]))
    wb.close()
    return ids


def _style_header(ws, ncols):
    for col_idx in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"


def append_rows(path: Path, rows):
    if path.exists():
        wb = openpyxl.load_workbook(path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    if "Jobs" not in wb.sheetnames:
        ws = wb.create_sheet("Jobs")
        ws.append(COLUMNS)
        _style_header(ws, len(COLUMNS))
    else:
        ws = wb["Jobs"]

    for row in rows:
        ws.append([row.get(c) for c in COLUMNS])

    if "Run Log" not in wb.sheetnames:
        log = wb.create_sheet("Run Log")
        log.append(["run_timestamp_utc", "new_rows_added", "total_rows_after"])
        _style_header(log, 3)
    else:
        log = wb["Run Log"]

    total_rows_after = ws.max_row - 1
    log.append([
        datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"),
        len(rows),
        total_rows_after,
    ])
    wb.save(path)


# --------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------

def main():
    if curl_requests is None:
        print("curl_cffi is not installed. Run: pip install -r requirements.txt")
        sys.exit(1)

    output_path = Path(OUTPUT_XLSX)
    existing_ids = load_existing_ids(output_path)
    print(f"Loaded {len(existing_ids)} existing job IDs from {output_path}")

    new_rows = []
    consecutive_seen = 0

    # Plain TLS-fingerprint-impersonating HTTP, no browser, no CDP — see
    # this file's 2026-10-02 (part 2) docstring note for why Playwright was
    # dropped entirely.
    session = curl_requests.Session(impersonate=IMPERSONATE)

    headers, base_url = discover_api_credentials(session)
    if not headers or not base_url:
        print("  Could not find API credentials — the site's behavior may "
              "have changed (a different chunk layout, a renamed anon-key "
              "pattern), or it may be blocking this request the same way "
              "ethiopianreporterjobs.com blocks curl_cffi (see that file's "
              "docstring). Aborting this run without writing anything, "
              "rather than committing a tracker that would look like a "
              "legitimate 'zero new postings' day.")
        sys.exit(1)
    print(f"  Using API base URL: {base_url}")

    for page_num in range(MAX_PAGES):
        offset = page_num * PAGE_SIZE
        try:
            batch = fetch_jobs_page(session, base_url, headers, offset)
        except Exception as e:
            if offset == 0:
                # The very first fetch failing isn't "zero jobs" — it's
                # "we don't know," the same distinction ReporterJobs'
                # ListingPageBlocked makes for its own first page.
                print(f"  Offset 0: fetch failed ({e}) — aborting this run without "
                      f"writing anything, rather than committing a tracker that "
                      f"would look like a legitimate 'zero new postings' day.")
                sys.exit(1)
            print(f"  Offset {offset}: fetch failed ({e}) — stopping pagination "
                  f"(keeping the {len(new_rows)} new posting(s) already found).")
            break

        if not batch:
            print(f"  Offset {offset}: no more jobs — reached the end.")
            break

        page_new = 0
        for row in batch:
            job_id = row.get("id")
            if job_id in existing_ids:
                consecutive_seen += 1
            else:
                consecutive_seen = 0
                new_rows.append(structure_job(row))
                existing_ids.add(job_id)
                page_new += 1

        print(f"  Offset {offset}: {len(batch)} jobs, {page_new} new, "
              f"{consecutive_seen} consecutive already-seen.")

        if consecutive_seen >= STOP_AFTER_CONSECUTIVE_SEEN:
            print(f"  Hit {STOP_AFTER_CONSECUTIVE_SEEN} consecutive already-seen postings — "
                  f"caught up, stopping pagination.")
            break
        if len(batch) < PAGE_SIZE:
            break
        time.sleep(REQUEST_DELAY_SECONDS)

    if new_rows:
        append_rows(output_path, new_rows)
        print(f"Added {len(new_rows)} new postings. Workbook: {output_path}")
    else:
        append_rows(output_path, [])
        print("No new postings this run.")


if __name__ == "__main__":
    main()
